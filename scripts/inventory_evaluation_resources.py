"""Inventory selected local research tables without exporting transcript rows."""
from __future__ import annotations
import argparse,csv,json,sys
from collections import Counter,defaultdict
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from src.evaluation.common import digest,private_directory,read_jsonl,write_json
from src.evaluation.text import normalize

def inventory(root,baseline,output):
    inputs={normalize(r['text_raw']):r['utterance_id'] for r in read_jsonl(baseline/'tables/metrics.jsonl')}
    records=[]
    # Top-level research CSVs only; no recursive exploration of user directories.
    for path in sorted(root.glob('*.csv')):
        counts=Counter();ids=Counter();texts=Counter();matches=defaultdict(set);strategies=defaultdict(set)
        with path.open(encoding='utf-8-sig',newline='') as f:
            reader=csv.DictReader(f);columns=reader.fieldnames or []
            for row in reader:
                counts['rows']+=1
                if row.get('name'):ids[row['name']]+=1
                text=row.get('Transcript')
                if text:
                    key=normalize(text);texts[key]+=1
                    if key in inputs:matches['Transcript'].add(inputs[key])
                    for col in ('Type of code-switching','Direction of code-switches','Number of code-switches'):
                        if row.get(col):strategies[(key,col)].add(row[col])
        records.append({'path':str(path.resolve()),'sha256':digest(path.read_bytes()),'bytes':path.stat().st_size,
            'rows':counts['rows'],'columns':columns,'unique_names':len(ids),'duplicate_name_groups':sum(n>1 for n in ids.values()),
            'duplicate_normalized_transcript_groups':sum(n>1 for n in texts.values()),
            'conflicting_strategy_transcript_field_groups':sum(len(v)>1 for v in strategies.values()),
            'baseline_exact_text_ids':{k:sorted(v,key=int) for k,v in matches.items()},
            'join_limit':'Exact normalized text is a candidate join; duplicate text and speaker-name homographs require review. No row-order joins.'})
    workbooks=[]
    import openpyxl
    for path in sorted(root.glob('*.xlsx')):
        workbook=openpyxl.load_workbook(path,read_only=True,data_only=True);sheets=[]
        for sheet in workbook:
            rows=sheet.iter_rows(values_only=True);header=next(rows,());nonempty=0;matches=defaultdict(set)
            # Inspect unnamed columns too: the original transcript column is unnamed.
            text_columns=[i for i,h in enumerate(header) if h is None or 'transcript' in str(h).lower()]
            for row in rows:
                nonempty+=any(v is not None for v in row)
                for i in text_columns:
                    if i<len(row) and isinstance(row[i],str) and normalize(row[i]) in inputs:matches[str(i)+':'+str(header[i])].add(inputs[normalize(row[i])])
            sheets.append({'sheet':sheet.title,'formatted_rows':sheet.max_row,'nonempty_data_rows':nonempty,'columns':list(header),
                'baseline_exact_text_ids':{k:sorted(v,key=int) for k,v in matches.items()}})
        workbook.close();workbooks.append({'path':str(path.resolve()),'sha256':digest(path.read_bytes()),'sheets':sheets})
    notebooks=[]
    for path in sorted(root.glob('*.ipynb')):
        nb=json.loads(path.read_text(encoding='utf-8'));sources=[''.join(c.get('source',[])) for c in nb.get('cells',[]) if c.get('cell_type')=='code']
        notebooks.append({'path':str(path.resolve()),'sha256':digest(path.read_bytes()),'code_cells':len(sources),
            'method_leads':{term:sum(term.lower() in s.lower() for s in sources) for term in ('Transcript','LID_tags','speaking','syllable','align','break','merge')},
            'audit':'Read source cells only; notebook outputs not executed or exported'})
    write_json(private_directory(output)/'local_resource_inventory.json',{'tables':records,'workbooks':workbooks,'notebooks':notebooks,
        'audio_inventory':{p.name:sum(1 for _ in p.glob('*.wav')) for p in (root/'wavfiles_utt_level',root/'wavfiles_conv_level',root/'ML_audio') if p.is_dir()},
        'data_provenance':'Existing user-supplied local files; automated annotation provenance requires source verification. No files are treated as ground truth from names alone.'})

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--root',type=Path,required=True);p.add_argument('--baseline',type=Path,required=True);p.add_argument('--output',type=Path,required=True)
    a=p.parse_args();inventory(a.root,a.baseline,a.output)
